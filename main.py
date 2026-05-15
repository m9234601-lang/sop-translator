import streamlit as st
import google.generativeai as genai
from openpyxl import load_workbook
import io
import time

# 1. 페이지 설정
st.set_page_config(page_title="SOP 양식 유지 번역기", layout="wide")
st.title("🏭 제조 SOP 양식 유지 번역기 (Gemini AI)")

# 2. 사이드바 설정 (API 키 입력)
with st.sidebar:
    st.header("⚙️ 설정")
    # Secrets에 등록했거나 직접 입력한 키를 가져옵니다.
    api_key = st.secrets.get("GOOGLE_API_KEY", "")
    if not api_key:
        api_key = st.text_input("Gemini API Key를 입력하세요", type="password")
    
    if api_key:
        genai.configure(api_key=api_key)
        st.success("✅ API 연결 완료")
    
    st.divider()
    target_lang = st.selectbox("목표 언어", ["영어", "한국어", "중국어", "베트남어", "일본어"])

# 3. 파일 업로드 및 로직
uploaded_file = st.file_uploader("원본 엑셀 파일 업로드", type=["xlsx"])

if uploaded_file and api_key:
    if st.button("양식 유지 번역 시작"):
        try:
            # 엑셀 로드
            wb = load_workbook(uploaded_file)
            model = genai.GenerativeModel('gemini-1.5-flash')
            
            # 번역 대상 셀 수집
            all_cells = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and len(str(cell.value).strip()) > 0:
                            all_cells.append(cell)
            
            progress_bar = st.progress(0)
            status_text = st.empty()

            for i, cell in enumerate(all_cells):
                status_text.text(f"⏳ 번역 중: {i+1}/{len(all_cells)} 셀 처리 중...")
                try:
                    # Gemini에게 번역 요청
                    prompt = f"Translate the following manufacturing SOP text into {target_lang}. Keep technical terms accurate. Translate only the text: {cell.value}"
                    response = model.generate_content(prompt)
                    if response.text:
                        cell.value = response.text.strip()
                except:
                    continue # 오류 시 해당 셀은 원문 유지
                
                progress_bar.progress((i + 1) / len(all_cells))
                time.sleep(0.1) # 안정적인 API 호출을 위한 지연

            # 결과물 생성
            output = io.BytesIO()
            wb.save(output)
            
            st.success("🎉 번역이 완료되었습니다!")
            st.download_button(
                label="📥 번역된 파일 다운로드",
                data=output.getvalue(),
                file_name=f"translated_sop.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.error(f"오류 발생: {e}")
